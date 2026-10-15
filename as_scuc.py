"""
Grid-India daily AS/SCUC report: download PDFs, extract SRAS/TRAS
summary tables, and consolidate into one monthly Excel.
"""

import json
import re
import time
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
import pdfplumber
import requests
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import (
    GRID_HEADERS, CDN_BASE, AS_CACHE_FILE, AS_METADATA_FILE,
    financial_year, expected_days,
    clean_filename,
    as_month_folder, as_panel_path,
)

AS_API_URL = "https://webapi.grid-india.in/api/v1/file"

# Set these manually to skip auto-detection once you know the real values
REPORT_SOURCE = None   # e.g. "GRDW"
REPORT_TYPE   = None   # e.g. "DAILY_AS_SCUC_REPORT"

CANDIDATE_SOURCES = ["GRDW"]
CANDIDATE_TYPES = [
    "DAILY_AS_SCUC_REPORT",
    "DAILY_AS_REPORT",
    "AS_SCUC_REPORT",
    "DAILY_ANCILLARY_SERVICES_REPORT",
    "ANCILLARY_SERVICES_REPORT",
    "DAILY_ANCILLARY_SERVICE_REPORT",
    "DAILY_AS_SCUC",
    "SCUC_REPORT",
    "AS_REPORT",
]

AS_COLUMNS = [
    "Date",
    "SRAS Up", "SRAS Down", "SRAS Max", "SRAS Min", "SRAS Avg",
    "TRAS Up", "TRAS Down", "TRAS Max", "TRAS Min", "TRAS Avg",
]


# --- GENERIC HELPERS ---

def as_date_from_title(title: str):
    """AS/SCUC titles carry a DD-MM-YYYY date (4-digit year, hyphens) -
    different from PSP's DD.MM.YY dotted format."""
    m = re.search(r"(\d{2})-(\d{2})-(\d{4})", title)
    if not m:
        return None
    d, mo, y = m.groups()
    try:
        return date(int(y), int(mo), int(d))
    except ValueError:
        return None


def as_belongs_to_month(title: str, cal_year: int, cal_month: int) -> bool:
    d = as_date_from_title(title)
    return d is not None and d.year == cal_year and d.month == cal_month


# --- API ---

def as_fetch_api(source: str, rtype: str, fy: str, month_str: str) -> list:
    payload = {"_source": source, "_type": rtype, "_fileDate": fy, "_month": month_str}
    resp = requests.post(AS_API_URL, json=payload, headers=GRID_HEADERS, verify=False, timeout=60)
    resp.raise_for_status()
    return resp.json().get("retData", [])


def _load_cached_type():
    if AS_CACHE_FILE.exists():
        try:
            d = json.loads(AS_CACHE_FILE.read_text())
            return d.get("source"), d.get("type")
        except Exception:
            return None, None
    return None, None


def _save_cached_type(source, rtype):
    AS_CACHE_FILE.parent.mkdir(parents=True, exist_ok=True)
    AS_CACHE_FILE.write_text(json.dumps({"source": source, "type": rtype}))


def detect_report_type(fy: str, month_str: str, force: bool = False):
    """Auto-detects the correct (_source, _type) pair by probing candidates
    against the live API and checking that returned titles actually look
    like AS_SCUC reports. Caches the result after first success."""
    if REPORT_SOURCE and REPORT_TYPE:
        return REPORT_SOURCE, REPORT_TYPE

    if not force:
        s, t = _load_cached_type()
        if s and t:
            return s, t

    print("  Auto-detecting AS/SCUC report type from the API (first run only)...")
    for source in CANDIDATE_SOURCES:
        for rtype in CANDIDATE_TYPES:
            try:
                records = as_fetch_api(source, rtype, fy, month_str)
            except Exception:
                continue
            if records and any(re.search(r"AS[_\s-]*SCUC", str(r.get("Title_", "")), re.I) for r in records):
                print(f"  Detected: _source='{source}', _type='{rtype}'  ({len(records)} records found)")
                _save_cached_type(source, rtype)
                return source, rtype
            time.sleep(0.3)

    raise RuntimeError(
        "Could not auto-detect the AS/SCUC report type from the API.\n"
        "Open browser DevTools -> Network tab on "
        "https://grid-india.in/en/reports/daily-as-report, find the POST request "
        "to webapi.grid-india.in/api/v1/file, and read its real _source/_type "
        "payload - then set REPORT_SOURCE / REPORT_TYPE at the top of this script."
    )


# --- DOWNLOADER ---

def fetch_as_scuc_month(cal_year: int, cal_month: int):
    month_name = date(cal_year, cal_month, 1).strftime("%B")
    month_padded = str(cal_month).zfill(2)
    fy = financial_year(cal_year, cal_month)
    days_in_month = expected_days(cal_year, cal_month)
    save_folder = as_month_folder(cal_year, cal_month)
    save_folder.mkdir(parents=True, exist_ok=True)

    print(f"\n{'='*60}")
    print(f"  Fetching Daily AS/SCUC Report: {month_name} {cal_year}  |  FY {fy}")
    print(f"  Expected days in month: {days_in_month}")
    print(f"  Save folder: {save_folder}")
    print(f"{'='*60}")

    source, rtype = detect_report_type(fy, month_padded)

    raw_records = as_fetch_api(source, rtype, fy, month_padded)
    records = [r for r in raw_records if as_belongs_to_month(r.get("Title_", ""), cal_year, cal_month)]
    print(f"  API returned {len(raw_records)} raw record(s), {len(records)} belong to {month_name} {cal_year}")

    # last-day spillover (same approach as the PSP script - some reports for the final day
    # of a month are only listed under the next month's API response)
    next_date = date(cal_year, cal_month, days_in_month) + timedelta(days=1)
    next_fy = financial_year(next_date.year, next_date.month)
    next_month_str = str(next_date.month).zfill(2)
    next_raw = as_fetch_api(source, rtype, next_fy, next_month_str)
    last_day = date(cal_year, cal_month, days_in_month)
    last_day_records = [r for r in next_raw if as_date_from_title(r.get("Title_", "")) == last_day]
    if last_day_records:
        print(f"  Found {len(last_day_records)} last-day record(s) via next-month lookup")

    all_records = records + last_day_records

    downloaded = skipped = failed = 0
    saved_records = []

    for item in all_records:
        mime = item.get("MimeType", "")
        if "pdf" not in mime.lower():
            continue

        title = item.get("Title_", "")
        file_path = item.get("FilePath", "")
        clean_title = re.sub(r"\.pdf$", "", clean_filename(title), flags=re.I)
        filename = clean_title + ".pdf"
        file_url = CDN_BASE + file_path
        save_path = save_folder / filename

        item["download_url"] = file_url
        item["local_file"] = str(save_path)
        saved_records.append(item)

        if save_path.exists():
            print(f"  [SKIP] {filename}")
            skipped += 1
            continue

        print(f"  [DOWN] {filename}")
        try:
            r = requests.get(file_url, headers=GRID_HEADERS, verify=False, timeout=120)
            if r.status_code == 200:
                save_path.write_bytes(r.content)
                downloaded += 1
            else:
                print(f"  [FAIL] HTTP {r.status_code} -> {file_url}")
                failed += 1
            time.sleep(1)
        except Exception as e:
            print(f"  [ERR]  {e}")
            failed += 1

    total = downloaded + skipped
    print(f"\n  Downloaded : {downloaded}")
    print(f"  Skipped    : {skipped}  (already existed)")
    print(f"  Failed     : {failed}")
    print(f"  Total PDFs : {total}  |  Expected: {days_in_month}")

    if total < days_in_month:
        print(f"  WARNING: {days_in_month - total} day(s) missing - may not yet be published.")
    elif total == days_in_month:
        print(f"  OK: All {days_in_month} daily reports accounted for.")

    if saved_records:
        df_new = pd.DataFrame(saved_records)
        if AS_METADATA_FILE.exists():
            df_existing = pd.read_csv(AS_METADATA_FILE)
            df_combined = pd.concat([df_existing, df_new], ignore_index=True).drop_duplicates(
                subset=["Title_", "local_file"], keep="last"
            )
        else:
            df_combined = df_new
        df_combined.to_csv(AS_METADATA_FILE, index=False)

    return saved_records


# --- PDF PARSER ---

def extract_date_from_as_filename(fname: str):
    m = re.search(r"(\d{2})-(\d{2})-(\d{4})", fname)
    return f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else None


def extract_as_scuc_pdf(filepath: Path) -> dict:
    """
    Pulls the SRAS/TRAS summary table (Up MWh, Down MWh, Max MW, Min MW,
    Avg MW per row) from the report. Regex-based rather than coordinate-based
    since every SRAS/TRAS row always carries all 5 values (no blank-cell risk
    like the NPP DGR16 report had). Negative lookbehind guards against the
    "SRAS+TRAS:" row being mistaken for a plain "TRAS:" match.
    """
    date_str = extract_date_from_as_filename(filepath.name)

    with pdfplumber.open(filepath) as pdf:
        full_text = "\n".join((p.extract_text() or "") for p in pdf.pages)

    def get_row(label):
        pattern = (
            rf"(?<![+\w]){label}\s*:\s*"
            r"(-?[\d,]+\.?\d*)\s+(-?[\d,]+\.?\d*)\s+(-?[\d,]+\.?\d*)\s+"
            r"(-?[\d,]+\.?\d*)\s+(-?[\d,]+\.?\d*)"
        )
        m = re.search(pattern, full_text)
        if not m:
            return [None] * 5
        return [float(g.replace(",", "")) for g in m.groups()]

    sras = get_row("SRAS")
    tras = get_row("TRAS")

    return {
        "Date": date_str,
        "SRAS Up": sras[0], "SRAS Down": sras[1], "SRAS Max": sras[2], "SRAS Min": sras[3], "SRAS Avg": sras[4],
        "TRAS Up": tras[0], "TRAS Down": tras[1], "TRAS Max": tras[2], "TRAS Min": tras[3], "TRAS Avg": tras[4],
    }


# --- MONTHLY EXCEL BUILDER ---

def build_as_scuc_month_excel(cal_year: int, cal_month: int, out_path=None):
    month_name = date(cal_year, cal_month, 1).strftime("%B")
    folder = as_month_folder(cal_year, cal_month)

    if not folder.exists():
        print(f"No folder found at {folder} - run fetch_as_scuc_month({cal_year}, {cal_month}) first.")
        return None

    pdf_files = sorted(folder.glob("*.pdf"))
    if not pdf_files:
        print(f"No PDF files found in {folder}.")
        return None

    print(f"\nFound {len(pdf_files)} file(s) in {folder}. Processing...")

    records, errors = [], []
    for f in pdf_files:
        try:
            rec = extract_as_scuc_pdf(f)
            if rec["Date"] is None:
                raise ValueError("could not parse date from filename")
            if None in (rec["SRAS Up"], rec["TRAS Up"]):
                print(f"  WARNING: {f.name}: SRAS/TRAS table not found - check PDF layout")
            records.append(rec)
            print(f"  OK: {f.name}  [{rec['Date']}]")
        except Exception as e:
            print(f"  ERROR: {f.name}: {e}")
            errors.append((f.name, str(e)))

    if not records:
        print("No data extracted. Exiting.")
        return None

    df = pd.DataFrame(records, columns=AS_COLUMNS)
    df["Date"] = pd.to_datetime(df["Date"], format="%d-%m-%Y")
    df = df.sort_values("Date").reset_index(drop=True)

    if out_path is None:
        out_path = as_panel_path(cal_year, cal_month)
    out_path = Path(out_path)

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="AS_SCUC_Monthly", index=False)
        ws = writer.sheets["AS_SCUC_Monthly"]

        date_fill  = "1F3864"
        sras_fill  = "2E75B6"
        tras_fill  = "C00000"
        thin = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for cell in ws[1]:
            val = str(cell.value or "")
            fill_hex = date_fill if val == "Date" else (sras_fill if val.startswith("SRAS") else tras_fill)
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill("solid", fgColor=fill_hex)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if cell.column == 1:
                    cell.number_format = "dd-mm-yyyy"
                elif isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0"

        ws.column_dimensions["A"].width = 14
        for i in range(2, len(AS_COLUMNS) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 12
        ws.freeze_panes = "B2"

        if errors:
            pd.DataFrame(errors, columns=["File", "Error"]).to_excel(writer, sheet_name="Errors", index=False)

    # missing-days report
    days_in_month = expected_days(cal_year, cal_month)
    present_days = set(df["Date"].dt.day)
    missing = sorted(set(range(1, days_in_month + 1)) - present_days)

    print(f"\n{'='*60}")
    print(f"  Month     : {month_name} {cal_year}  ({days_in_month} days)")
    print(f"  Extracted : {len(df)} days")
    if missing:
        print(f"  Missing   : {len(missing)} day(s) -> {missing}")
    else:
        print(f"  Complete  : all {days_in_month} days present")
    if errors:
        print(f"  Failed    : {len(errors)} file(s) - see 'Errors' sheet")
    print(f"  Saved     : {out_path}")
    print(f"{'='*60}")

    return df


# --- ORCHESTRATOR ---

def download_and_build_as_scuc_month(cal_year: int, cal_month: int):
    fetch_as_scuc_month(cal_year, cal_month)
    return build_as_scuc_month_excel(cal_year, cal_month)
