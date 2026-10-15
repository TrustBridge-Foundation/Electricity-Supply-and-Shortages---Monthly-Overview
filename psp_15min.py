"""
Extract the 15-minute SCADA data (sheet 'TimeSeries') of the two extreme days
of a month from the daily Grid-India PSP files.

  Peak day   = day of HIGHEST all-India electricity requirement in the month
  Lowest day = day of LOWEST  all-India electricity requirement in the month

with requirement = Energy Met + Energy Shortage (MU), as in plot_energy_met().
The two days are chosen from the monthly PSP time-series workbook, their daily
PSP files are located in the month's PSP folder, and the 96 blocks of each are
consolidated into one Excel file:

    GRID_PSP_REPORTS/Extreme_Days_15Min/<FY>/extreme_days_15min_<month><year>.xlsx
        Sheet 'Day_Selection' : which days were picked and why
        Sheet '15Min_Data'    : 2 x 96 rows, one column per series

The output deliberately lives OUTSIDE the monthly PSP folder: build_monthly_
timeseries() treats every .xls/.xlsx file in that folder as a daily report.

Usage:
    from psp_15min import build_extreme_days_15min
    build_extreme_days_15min(2026, 9)
"""

import re
from dataclasses import dataclass
from datetime import date, datetime, time as dtime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import (
    PSP_ROOT, financial_year, expected_days,
    psp_month_dir, psp_timeseries_path,
)

# --- Settings ---

COL_MET   = 'A | Energy Met All-India (MU)'
COL_SHORT = 'A | Energy Shortage All-India (MU)'

# "requirement" -> Energy Met + Energy Shortage (default);  "met" -> Energy Met only
DAY_SELECTION = "requirement"

EXTREME_ROOT = PSP_ROOT / "Extreme_Days_15Min"

# TimeSeries header (normalised prefix) -> output column name. Order matters
# only in that each sheet column is claimed by the first prefix that matches.
HEADER_MAP = [
    ("FREQUENCY",        "Frequency (Hz)"),
    ("DEMAND MET",       "Demand Met (MW)"),
    ("STORAGE DEMAND",   "Storage Demand (MW)"),
    ("NUCLEAR",          "Nuclear (MW)"),
    ("WIND",             "Wind (MW)"),
    ("SOLAR",            "Solar (MW)"),
    ("HYDRO",            "Hydro (MW)"),
    ("GAS",              "Gas (MW)"),
    ("THERMAL",          "Thermal (MW)"),
    ("STORAGE (",        "Storage Discharge (MW)"),
    ("OTHERS",           "Others (MW)"),
    ("NET DEMAND MET",   "Net Demand Met (MW)"),
    ("TOTAL GENERATION", "Total Generation (MW)"),
    ("NET TRANSNATIONAL", "Net Transnational Exchange (MW)"),
]
DATA_COLUMNS = [name for _, name in HEADER_MAP]
STORAGE_COLUMNS = {"Storage Demand (MW)", "Storage Discharge (MW)"}  # blank = no activity -> 0

# Columns the analysis can do without. Earlier PSP files (e.g. April 2026) have no
# storage columns: those are left blank (NaN, i.e. not reported), and net demand and
# total generation are derived from the other columns.
OPTIONAL_COLUMNS = {"Storage Demand (MW)", "Storage Discharge (MW)", "Net Demand Met (MW)",
                    "Total Generation (MW)", "Net Transnational Exchange (MW)"}
SOURCE_COLUMNS = ["Nuclear (MW)", "Wind (MW)", "Solar (MW)", "Hydro (MW)", "Gas (MW)",
                  "Thermal (MW)", "Others (MW)", "Storage Discharge (MW)"]

OUT_COLUMNS = ["Date", "Day Type", "Block", "Time", "Hour"] + DATA_COLUMNS


# --- Day selection ---

def ordinal(n: int) -> str:
    suffix = 'th' if 11 <= n % 100 <= 13 else {1: 'st', 2: 'nd', 3: 'rd'}.get(n % 10, 'th')
    return f"{n}{suffix}"


@dataclass(frozen=True)
class DayPick:
    day: date
    value_mu: float
    kind: str                     # "peak" or "lowest"

    @property
    def label(self) -> str:       # e.g. "9th September"
        return f"{ordinal(self.day.day)} {self.day.strftime('%B')}"

    @property
    def panel_title(self) -> str:
        return f"{self.label} ({self.kind} requirement)"


def select_extreme_days(psp_path, cal_year, cal_month, metric=DAY_SELECTION):
    """Return (peak_day, lowest_day) as DayPick objects from the monthly PSP file."""
    df = pd.read_excel(psp_path)
    df['Date'] = pd.to_datetime(df['Date'], format='%d-%m-%Y')
    met   = pd.to_numeric(df[COL_MET],   errors='coerce')
    short = pd.to_numeric(df[COL_SHORT], errors='coerce')

    if metric == "requirement":
        series = met + short
    elif metric == "met":
        series = met
    else:
        raise ValueError("DAY_SELECTION must be 'requirement' or 'met'.")

    valid = series.notna()
    series, dates = series[valid], df.loc[valid, 'Date']
    if len(series) < 2:
        raise ValueError(f"Need at least two valid days in {psp_path}; found {len(series)}.")

    total = expected_days(cal_year, cal_month)
    if len(series) < total:
        print(f"  WARNING: Only {len(series)} of {total} days have valid PSP data - "
              f"the selected extremes are provisional until the month is complete.")

    i_max, i_min = series.idxmax(), series.idxmin()
    peak   = DayPick(dates.loc[i_max].date(), float(series.loc[i_max]), "peak")
    lowest = DayPick(dates.loc[i_min].date(), float(series.loc[i_min]), "lowest")

    print(f"\n  Day selection by '{metric}' over {len(series)} days "
          f"(month mean {series.mean():,.0f} MU)")
    print(f"    Highest : {peak.day:%d-%m-%Y}  {peak.value_mu:,.0f} MU")
    print(f"    Lowest  : {lowest.day:%d-%m-%Y}  {lowest.value_mu:,.0f} MU")
    return peak, lowest


# --- Locate a day's PSP file ---

_FILE_DATE = re.compile(r"^(\d{2})[._](\d{2})[._](\d{2})")


def find_daily_psp_file(cal_year, cal_month, day: date) -> Path:
    """File names start with DD.MM.YY (e.g. '09.09.26_NLDC_PSP.xls')."""
    folder = psp_month_dir(cal_year, cal_month)
    if not folder.exists():
        raise FileNotFoundError(f"PSP folder not found: {folder}")

    matches = []
    for f in sorted(folder.iterdir()):
        if f.suffix.lower() not in (".xls", ".xlsx") or f.name.startswith(("~", "monthly_timeseries")):
            continue
        m = _FILE_DATE.match(f.name)
        if m and (int(m.group(1)), int(m.group(2)), 2000 + int(m.group(3))) == (day.day, day.month, day.year):
            matches.append(f)

    if not matches:
        raise FileNotFoundError(f"No PSP file for {day:%d-%m-%Y} in {folder}")
    if len(matches) > 1:
        print(f"  WARNING: {len(matches)} files for {day:%d-%m-%Y}: {[f.name for f in matches]} "
              f"- using the last one by name ({matches[-1].name}).")
    return matches[-1]


# --- Read the TimeSeries sheet ---

def _read_sheet_rows(path: Path):
    """Return the TimeSeries sheet as a list of row tuples (.xls or .xlsx)."""
    is_ts = lambda n: n.strip().lower().replace(" ", "").replace("_", "") == "timeseries"

    if path.suffix.lower() == ".xls":
        import xlrd
        wb = xlrd.open_workbook(str(path))
        name = next((n for n in wb.sheet_names() if is_ts(n)), None)
        if name is None:
            raise ValueError(f"No 'TimeSeries' sheet in {path.name}. Sheets: {wb.sheet_names()}")
        sh = wb.sheet_by_name(name)
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == 1:
                    row.append(cell.value if cell.value.strip() else None)
                elif cell.ctype == 2:
                    row.append(cell.value)
                elif cell.ctype == 3:
                    row.append(xlrd.xldate_as_datetime(cell.value, wb.datemode))
                else:
                    row.append(None)
            rows.append(tuple(row))
        return rows

    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        name = next((n for n in wb.sheetnames if is_ts(n)), None)
        if name is None:
            raise ValueError(f"No 'TimeSeries' sheet in {path.name}. Sheets: {wb.sheetnames}")
        return list(wb[name].iter_rows(values_only=True))
    finally:
        wb.close()


def _norm_header(cell) -> str:
    """Upper-case, strip footnote superscripts and collapse whitespace/line breaks."""
    text = re.sub(r"[\u00b9\u00b2\u00b3\u2070-\u2079]", "", str(cell))
    return " ".join(text.split()).upper()


def _minutes(value):
    """TIME cell -> minutes since midnight, or None if it is not a time."""
    if isinstance(value, datetime):
        return round((value.hour * 3600 + value.minute * 60 + value.second) / 60)
    if isinstance(value, dtime):
        return round((value.hour * 3600 + value.minute * 60 + value.second) / 60)
    if isinstance(value, str):
        m = re.fullmatch(r"\s*(\d{1,2}):(\d{2})(?::\d{2})?\s*", value)
        return int(m.group(1)) * 60 + int(m.group(2)) if m else None
    if isinstance(value, float) and 0 <= value < 1:        # Excel time fraction
        return round(value * 24 * 60)
    return None


def _num(value):
    if isinstance(value, bool) or value is None:
        return np.nan
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return np.nan


def read_timeseries_15min(path, day: date) -> pd.DataFrame:
    """Return the day's 15-minute data as a tidy DataFrame (one row per block).
    The date comes from the file name: 'Date of Reporting' inside the sheet is
    the following day."""
    path = Path(path)
    rows = _read_sheet_rows(path)

    required = [name for _, name in HEADER_MAP if name not in OPTIONAL_COLUMNS]
    header_idx, time_col, cols = None, None, {}
    best = (0, None, {}, [])           # closest header row, kept for the error message
    for i, row in enumerate(rows[:30]):
        norms = {c: _norm_header(v) for c, v in enumerate(row) if v is not None}
        t_cols = [c for c, n in norms.items() if n == "TIME" or n.startswith("TIME ")]
        found, claimed = {}, set()
        for prefix, out_name in HEADER_MAP:
            for c, n in norms.items():
                if c not in claimed and n.startswith(prefix):
                    found[out_name] = c
                    claimed.add(c)
                    break
        if t_cols and all(name in found for name in required):
            header_idx, time_col, cols = i, t_cols[0], found
            break
        score = len(found) + (1 if t_cols else 0)
        if score > best[0]:
            best = (score, i, found, list(norms.values()))

    if header_idx is None:
        _, i, found, cells = best
        if i is None:
            cells = [" ".join(str(v).split()) for r in rows[:6] for v in r if v is not None][:20]
        missing = [name for name in required if name not in found]
        raise ValueError(
            f"{path.name}: could not find the TIME header row with all required columns "
            f"in 'TimeSeries'.\n"
            f"  Required columns not found: {missing}\n"
            f"  Closest header row (sheet row {i + 1 if i is not None else 'n/a'}): {cells}")

    records = []
    for row in rows[header_idx + 1:]:
        mins = _minutes(row[time_col] if time_col < len(row) else None)
        if mins is None:
            continue                                   # sub-header / footnote rows
        rec = {"_min": mins}
        for out_name, c in cols.items():
            v = _num(row[c] if c < len(row) else None)
            rec[out_name] = 0.0 if (np.isnan(v) and out_name in STORAGE_COLUMNS) else v
        records.append(rec)

    df = pd.DataFrame(records).sort_values("_min").reset_index(drop=True)
    present = list(cols)

    absent = [name for name in DATA_COLUMNS if name not in cols]
    if absent:
        if "Storage Demand (MW)" in absent:
            df["Storage Demand (MW)"] = np.nan
        if "Storage Discharge (MW)" in absent:
            df["Storage Discharge (MW)"] = np.nan
        if "Net Demand Met (MW)" in absent:
            df["Net Demand Met (MW)"] = df["Demand Met (MW)"] - df["Wind (MW)"] - df["Solar (MW)"]
        if "Total Generation (MW)" in absent:
            df["Total Generation (MW)"] = df[SOURCE_COLUMNS].sum(axis=1)
        if "Net Transnational Exchange (MW)" in absent:
            df["Net Transnational Exchange (MW)"] = np.nan
        print(f"  NOTE: {path.name}: older layout, not in this file: {absent}")
    df = df[["_min"] + DATA_COLUMNS]

    if len(df) != 96:
        print(f"  WARNING: {path.name}: {len(df)} blocks read instead of 96.")
    if df["_min"].duplicated().any():
        print(f"  WARNING: {path.name}: duplicate time stamps found.")
    n_nan = int(df[present].isna().sum().sum())
    if n_nan:
        print(f"  WARNING: {path.name}: {n_nan} missing value(s) in the 15-min data.")

    df.insert(0, "Date", day)
    df.insert(1, "Block", df["_min"] // 15 + 1)
    df.insert(2, "Time", df["_min"].map(lambda m: f"{m // 60:02d}:{m % 60:02d}"))
    df.insert(3, "Hour", df["_min"] / 60)
    return df.drop(columns="_min")


# --- Output ---

def extreme_days_15min_path(cal_year: int, cal_month: int) -> Path:
    month_name = date(cal_year, cal_month, 1).strftime("%B").lower()
    return (EXTREME_ROOT / financial_year(cal_year, cal_month)
            / f"extreme_days_15min_{month_name}{cal_year}.xlsx")


def _format_sheet(ws, date_cols=(), width=14):
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if cell.column in date_cols:
                cell.number_format = "dd-mm-yyyy"
    for i in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"


def build_extreme_days_15min(cal_year: int, cal_month: int, out_path=None):
    """Pick the peak / lowest requirement days of the month, pull their 15-min
    TimeSeries data and write the consolidated workbook. Returns its path.
    Always rebuilt (it is small), so the selection stays in step with the PSP file."""
    month_name = date(cal_year, cal_month, 1).strftime("%B")
    print(f"\n{'='*60}\n  15-min data for extreme-requirement days: {month_name} {cal_year}\n{'='*60}")

    psp_path = psp_timeseries_path(cal_year, cal_month)
    if not psp_path.exists():
        raise FileNotFoundError(f"Monthly PSP workbook not found: {psp_path} "
                                f"- run build_monthly_timeseries first.")

    picks = select_extreme_days(psp_path, cal_year, cal_month)

    frames, selection = [], []
    for pick in picks:
        f = find_daily_psp_file(cal_year, cal_month, pick.day)
        print(f"  Reading {pick.kind:<6} day {pick.day:%d-%m-%Y}: {f.name}")
        df = read_timeseries_15min(f, pick.day)
        df.insert(1, "Day Type", pick.kind.capitalize())
        frames.append(df)
        selection.append({
            "Date": pick.day, "Day Type": pick.kind.capitalize(),
            "Selection metric": DAY_SELECTION, "Value (MU)": round(pick.value_mu, 2),
            "Source file": f.name,
        })

    data = pd.concat(frames, ignore_index=True)[OUT_COLUMNS]

    out_path = Path(out_path) if out_path else extreme_days_15min_path(cal_year, cal_month)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        pd.DataFrame(selection).to_excel(writer, sheet_name="Day_Selection", index=False)
        data.to_excel(writer, sheet_name="15Min_Data", index=False)
        _format_sheet(writer.sheets["Day_Selection"], date_cols=(1,), width=18)
        _format_sheet(writer.sheets["15Min_Data"], date_cols=(1,), width=15)
        writer.sheets["Day_Selection"].column_dimensions["E"].width = 30

    print(f"  OK: {len(data)} rows written -> {out_path}")
    return out_path
