"""
Grid-India daily PSP report: download monthly files and build a
consolidated monthly time-series Excel.
"""

import calendar
import re
import time
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
import requests
import xlrd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


from config import (
    GRID_HEADERS, CDN_BASE, PSP_METADATA_FILE,
    financial_year, expected_days,
    clean_filename, to_num,
    psp_month_dir,
)


PSP_API_URL  = "https://webapi.grid-india.in/api/v1/file"

def strip_known_extensions(name: str) -> str:
    return re.sub(r'\.(xlsx|xls)$', '', name, flags=re.IGNORECASE)

def psp_fetch_api(fy: str, month_str: str) -> list:
    payload = {
        "_source":   "GRDW",
        "_type":     "DAILY_PSP_REPORT",
        "_fileDate": fy,
        "_month":    month_str
    }
    resp = requests.post(PSP_API_URL, json=payload, headers=GRID_HEADERS, verify=False, timeout=60)
    resp.raise_for_status()
    return resp.json().get("retData", [])

def psp_belongs_to_month(title: str, month_padded: str, year_2d: str) -> bool:
    """
    Return True only if the title's date portion matches DD.<month_padded>.<year_2d>.
    Title format expected: DD.MM.YY_...  e.g. '31.01.26_NLDC_PSP'
    """
    pattern = rf"^\d{{2}}\.{re.escape(month_padded)}\.{re.escape(year_2d)}"
    return bool(re.match(pattern, title))

# --- main function ---

def fetch_psp_month(cal_year: int, cal_month: int):
    """
    Download all Grid India NLDC PSP Excel files for a given calendar month.
    Filters out any spillover records from adjacent months the API may return.
    """
    month_name    = date(cal_year, cal_month, 1).strftime("%B")   # "January"
    month_padded  = str(cal_month).zfill(2)                       # "01"
    year_2d       = str(cal_year)[-2:]                            # "26"
    fy            = financial_year(cal_year, cal_month)
    days_in_month = expected_days(cal_year, cal_month)

    save_folder = psp_month_dir(cal_year, cal_month)
    save_folder.mkdir(parents=True, exist_ok=True)

    print(f"\n{'='*60}")
    print(f"  Fetching PSP: {month_name} {cal_year}  |  FY {fy}")
    print(f"  Expected days in month: {days_in_month}")
    print(f"  Save folder: {save_folder}")
    print(f"{'='*60}")

    # --- Step 1: fetch primary month records & FILTER to target month ---
    raw_records = psp_fetch_api(fy, month_padded)
    print(f"  API returned {len(raw_records)} raw records for month {month_padded}")

    records = [
        r for r in raw_records
        if psp_belongs_to_month(r.get("Title_", ""), month_padded, year_2d)
    ]
    print(f"  After month filter: {len(records)} records belong to {month_name} {cal_year}")

    spillover = len(raw_records) - len(records)
    if spillover:
        print(f"  WARNING: Dropped {spillover} spillover record(s) from adjacent months")

    # --- Step 2: fetch Day-1 of next month to capture last-day report ---
    next_date      = date(cal_year, cal_month, days_in_month) + timedelta(days=1)
    next_fy        = financial_year(next_date.year, next_date.month)
    next_month_str = str(next_date.month).zfill(2)
    last_day_prefix = f"{str(days_in_month).zfill(2)}.{month_padded}.{year_2d}"

    print(f"  Fetching next-month Day-1 ({next_date}) to catch last-day report...")
    next_raw = psp_fetch_api(next_fy, next_month_str)

    last_day_records = [
        r for r in next_raw
        if r.get("Title_", "").startswith(last_day_prefix)
    ]
    print(f"  Found {len(last_day_records)} last-day record(s) matching '{last_day_prefix}*'")

    all_records = records + last_day_records

    # --- Step 3: download Excel files ---
    downloaded, skipped, failed = 0, 0, 0
    saved_records = []

    for item in all_records:
        mime      = item.get("MimeType", "")
        file_path = item.get("FilePath", "")
        title     = item.get("Title_", "")

        if not any(kw in mime.lower() for kw in ("excel", "spreadsheet", "xls")):
            continue

        clean_title = strip_known_extensions(clean_filename(title))
        ext         = ".xlsx" if file_path.lower().endswith(".xlsx") else ".xls"
        filename    = clean_title + ext
        file_url    = CDN_BASE + file_path
        save_path   = save_folder / filename

        item["download_url"] = file_url
        item["local_file"]   = str(save_path)
        saved_records.append(item)

        if save_path.exists():
            print(f"  [SKIP] {filename}")
            skipped += 1
            continue

        print(f"  [DOWN] {filename}")
        try:
            r = requests.get(file_url, headers=GRID_HEADERS, verify=False, timeout=120)
            if r.status_code == 200:
                save_path.write_bytes(r.content)
                downloaded += 1
            else:
                print(f"  [FAIL] HTTP {r.status_code} -> {file_url}")
                failed += 1
            time.sleep(1)
        except Exception as e:
            print(f"  [ERR]  {e}")
            failed += 1

    # --- Step 4: completeness check ---
    total_excel = downloaded + skipped
    print(f"\n  Downloaded : {downloaded}")
    print(f"  Skipped    : {skipped}  (already existed)")
    print(f"  Failed     : {failed}")
    print(f"  Total Excel: {total_excel}  |  Expected: {days_in_month}")

    if total_excel < days_in_month:
        print(f"  WARNING: {days_in_month - total_excel} day(s) missing - "
              f"may not yet be published on the portal.")
    elif total_excel == days_in_month:
        print(f"  OK: All {days_in_month} daily reports accounted for.")
    else:
        print("  NOTE: More files than days - check for duplicates or revised reports.")

    # --- Step 5: update metadata CSV ---
    if saved_records:
        df_new = pd.DataFrame(saved_records)
        if PSP_METADATA_FILE.exists():
            df_existing = pd.read_csv(PSP_METADATA_FILE)
            df_combined = pd.concat([df_existing, df_new], ignore_index=True).drop_duplicates()
        else:
            df_combined = df_new
        df_combined.to_csv(PSP_METADATA_FILE, index=False)

    return saved_records


def find_row(rows, col0_contains, after=0):
    """
    Return the index of the first row whose column-0 contains the search string
    (case-insensitive), starting from `after`. Returns -1 if not found.
    """
    needle = col0_contains.lower()
    for i in range(after, len(rows)):
        cell = rows[i][0]
        if cell:
            cell_normalized = " ".join(str(cell).lower().split())
            if needle in cell_normalized:
                return i

    return -1


def date_from_filename(filename):
    """
    Parse DD.MM.YY from filename and return DD-MM-YYYY string.
    e.g. '01.01.26_NLDC_PSP.xls'  ->  '01-01-2026'
         '09.06.26_NLDC_PSP_128.xls'  ->  '09-06-2026'
    """
    m = re.search(r'(\d{2})\.(\d{2})\.(\d{2})', filename)
    if m:
        d, mo, y = m.groups()
        return datetime.strptime(f"{d}-{mo}-20{y}", "%d-%m-%Y").strftime("%d-%m-%Y")
    raise ValueError(f"Cannot parse date from filename: {filename}")


def read_xls_rows(file_path):
    """
    Read all rows of the MOP_E sheet from a .xls file using xlrd.
    Returns a list of tuples (one per row), matching openpyxl's values_only format.
    """
    wb   = xlrd.open_workbook(str(file_path))

    # Find MOP_E sheet (handle minor name variations)
    sheet = None
    for name in wb.sheet_names():
        if "mop" in name.lower():
            sheet = wb.sheet_by_name(name)
            break
    if sheet is None:
        raise ValueError(
            f"MOP_E sheet not found in {file_path.name}. "
            f"Available: {wb.sheet_names()}"
        )

    rows = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            # xlrd cell types: 0=empty, 1=text, 2=number, 3=date, 4=bool, 5=error
            if cell.ctype == 0:
                row.append(None)
            elif cell.ctype == 1:
                row.append(cell.value if cell.value.strip() else None)
            elif cell.ctype == 2:
                # Return int if whole number, else float
                v = cell.value
                row.append(int(v) if v == int(v) else v)
            elif cell.ctype == 3:
                # Date - return as string HH:MM for time values
                dt = xlrd.xldate_as_datetime(cell.value, wb.datemode)
                row.append(dt.strftime("%H:%M"))
            elif cell.ctype == 4:
                row.append(bool(cell.value))
            else:
                row.append(None)
        rows.append(tuple(row))

    return rows


def read_xlsx_rows(file_path):
    """
    Read all rows of the MOP_E sheet from a .xlsx file using openpyxl.
    Returns a list of tuples (values only).
    """
    wb = load_workbook(str(file_path), read_only=True, data_only=True)

    sheet = None
    for name in wb.sheetnames:
        if "mop" in name.lower():
            sheet = wb[name]
            break
    if sheet is None:
        raise ValueError(
            f"MOP_E sheet not found in {file_path.name}. "
            f"Available: {wb.sheetnames}"
        )

    rows = list(sheet.iter_rows(values_only=True))
    wb.close()
    return rows


def load_rows(file_path):
    """Dispatch to the right reader based on file extension."""
    file_path = Path(file_path)
    ext       = file_path.suffix.lower()

    if ext == ".xls":
        try:
            return read_xls_rows(file_path)
        except ImportError:
            raise ImportError(
                "xlrd is required to read .xls files.\n"
                "Install it with:  pip install xlrd==1.2.0"
            )
    elif ext == ".xlsx":
        return read_xlsx_rows(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


# --- per-file extractor ---

def extract_one_file(file_path):
    """
    Extract all target variables from one daily PSP file.
    Returns a flat dict - one key per output column.
    Note: this will not work if the column numbers have been changed in the main pdf as we are extracting by column number. 
    """
    file_path = Path(file_path)
    rows      = load_rows(file_path)
    rec       = {}

    # --- Date from filename (actual report date, not 'Date of Reporting' inside) ---
    rec["Date"] = date_from_filename(file_path.name)

    # --- Table A: Energy Met & Energy Shortage - All-India TOTAL (col 7) ---
    r_em = find_row(rows, "Energy Met (MU)")
    r_es = find_row(rows, "Energy Shortage (MU)")
    rec["A_Energy_Met_AllIndia_MU"]      = to_num(rows[r_em][7]) if r_em >= 0 else None
    rec["A_Energy_Shortage_AllIndia_MU"] = to_num(rows[r_es][7]) if r_es >= 0 else None

    # --- Table F: Generation Outage - TOTAL col 7, %Share col 8 ---
    r_f       = find_row(rows, "F. Generation Outage")
    r_central = find_row(rows, "Central Sector", after=r_f) if r_f >= 0 else -1
    r_state   = find_row(rows, "State Sector",   after=r_f) if r_f >= 0 else -1

    # Find Total row strictly within Table F (next 10 rows after F header)
    r_f_total = -1
    if r_f >= 0:
        for i in range(r_f, min(r_f + 10, len(rows))):
            cell = rows[i][0]
            if cell and str(cell).strip().lower() == "total":
                r_f_total = i
                break

    rec["F_Central_Outage_AllIndia_MW"] = to_num(rows[r_central][7]) if r_central >= 0 else None
    rec["F_Central_Outage_PctShare"]    = to_num(rows[r_central][8]) if r_central >= 0 else None
    rec["F_State_Outage_AllIndia_MW"]   = to_num(rows[r_state][7])   if r_state   >= 0 else None
    rec["F_State_Outage_PctShare"]      = to_num(rows[r_state][8])   if r_state   >= 0 else None
    rec["F_Total_Outage_AllIndia_MW"]   = to_num(rows[r_f_total][7]) if r_f_total >= 0 else None

    # --- Table G: Sourcewise Generation - All India col 7, %Share col 8 ---
    r_g     = find_row(rows, "G. Sourcewise generation")
    after_g = r_g if r_g >= 0 else 0

    sources = {
        "Coal":    "Coal",
        "Lignite": "Lignite",
        "Hydro":   "Hydro",
        "Nuclear": "Nuclear",
        "Gas":     "Gas, Naptha",
        "RES":     "RES (Wind",
        "Total":   None,          # special: exact match "Total"
    }

    for label, search_str in sources.items():
        if label == "Total":
            ri = -1
            for i in range(after_g, min(after_g + 15, len(rows))):
                cell = rows[i][0]
                if cell and str(cell).strip().lower() == "total":
                    ri = i
                    break
        else:
            ri = find_row(rows, search_str, after=after_g)

        rec[f"G_{label}_AllIndia_MU"] = to_num(rows[ri][7]) if ri >= 0 else None
        rec[f"G_{label}_PctShare"]    = to_num(rows[ri][8]) if ri >= 0 else None

    r_res_share = find_row(rows, "Share of RES in total", after=after_g)
    r_nonfos    = find_row(rows, "Share of Non-fossil",   after=after_g)
    rec["G_RES_ShareOfTotal_Pct"]       = to_num(rows[r_res_share][7]) if r_res_share >= 0 else None
    rec["G_NonFossil_ShareOfTotal_Pct"] = to_num(rows[r_nonfos][7])   if r_nonfos    >= 0 else None

    # --- Table I: Solar & Non-Solar Peak Demand ---
    r_i = find_row(rows, "Based on State Max Demands")
    if r_i >= 0:
        solar    = rows[r_i]
        nonsolar = rows[r_i + 1] if r_i + 1 < len(rows) else (None,) * 16

        rec["I_Solar_MaxDemandMet_MW"]    = to_num(solar[5])
        rec["I_Solar_Time"]               = str(solar[7]).strip()    if solar[7]    else None
        rec["I_Solar_Shortage_MW"]        = to_num(solar[8])
        rec["I_NonSolar_MaxDemandMet_MW"] = to_num(nonsolar[5])
        rec["I_NonSolar_Time"]            = str(nonsolar[7]).strip() if nonsolar[7] else None
        rec["I_NonSolar_Shortage_MW"]     = to_num(nonsolar[8])
    else:
        for key in ["I_Solar_MaxDemandMet_MW", "I_Solar_Time", "I_Solar_Shortage_MW",
                    "I_NonSolar_MaxDemandMet_MW", "I_NonSolar_Time", "I_NonSolar_Shortage_MW"]:
            rec[key] = None

    return rec


# --- column rename map ---

RENAME_MAP = {
    "Date":                           "Date",
    "A_Energy_Met_AllIndia_MU":       "A | Energy Met All-India (MU)",
    "A_Energy_Shortage_AllIndia_MU":  "A | Energy Shortage All-India (MU)",
    "F_Central_Outage_AllIndia_MW":   "F | Central Sector Outage All-India (MW)",
    "F_Central_Outage_PctShare":      "F | Central Sector Outage % Share",
    "F_State_Outage_AllIndia_MW":     "F | State Sector Outage All-India (MW)",
    "F_State_Outage_PctShare":        "F | State Sector Outage % Share",
    "F_Total_Outage_AllIndia_MW":     "F | Total Outage All-India (MW)",
    "G_Coal_AllIndia_MU":             "G | Coal All-India (MU)",
    "G_Coal_PctShare":                "G | Coal % Share",
    "G_Lignite_AllIndia_MU":          "G | Lignite All-India (MU)",
    "G_Lignite_PctShare":             "G | Lignite % Share",
    "G_Hydro_AllIndia_MU":            "G | Hydro All-India (MU)",
    "G_Hydro_PctShare":               "G | Hydro % Share",
    "G_Nuclear_AllIndia_MU":          "G | Nuclear All-India (MU)",
    "G_Nuclear_PctShare":             "G | Nuclear % Share",
    "G_Gas_AllIndia_MU":              "G | Gas/Naphtha/Diesel All-India (MU)",
    "G_Gas_PctShare":                 "G | Gas/Naphtha/Diesel % Share",
    "G_RES_AllIndia_MU":              "G | RES All-India (MU)",
    "G_RES_PctShare":                 "G | RES % Share",
    "G_Total_AllIndia_MU":            "G | Total Generation All-India (MU)",
    "G_Total_PctShare":               "G | Total Generation % Share",
    "G_RES_ShareOfTotal_Pct":         "G | RES Share of Total Generation (%)",
    "G_NonFossil_ShareOfTotal_Pct":   "G | Non-Fossil Share of Total Generation (%)",
    "I_Solar_MaxDemandMet_MW":        "I | Solar-hr Max Demand Met (MW)",
    "I_Solar_Time":                   "I | Solar-hr Time of Max Demand",
    "I_Solar_Shortage_MW":            "I | Solar-hr Shortage (MW)",
    "I_NonSolar_MaxDemandMet_MW":     "I | Non-Solar-hr Max Demand Met (MW)",
    "I_NonSolar_Time":                "I | Non-Solar-hr Time of Max Demand",
    "I_NonSolar_Shortage_MW":         "I | Non-Solar-hr Shortage (MW)",
}

# --- main ---

def build_monthly_timeseries(folder, out_path=None):
    """
    Process all PSP Excel files in `folder` and write a consolidated
    monthly time-series Excel to `out_path` (defaults to inside `folder`).
    """
    folder = Path(folder)
    if not folder.exists():
        raise FileNotFoundError(f"Folder not found: {folder}")

    # Collect .xls and .xlsx files, skip temp/output files
    files = sorted(
        [f for f in folder.iterdir()
         if f.suffix.lower() in (".xls", ".xlsx")
         and not f.name.startswith("~")
         and not f.name.startswith("monthly_timeseries")],
        key=lambda f: f.name
    )

    if not files:
        raise FileNotFoundError(f"No PSP Excel files found in: {folder}")

    print(f"\n{'='*65}")
    print(f"  Folder : {folder}")
    print(f"  Files  : {len(files)} PSP files found")
    print(f"{'='*65}\n")

    records, errors = [], []

    for f in files:
        try:
            print(f"  Processing: {f.name} ...", end=" ", flush=True)
            rec = extract_one_file(f)
            records.append(rec)
            print(f"OK: [{rec['Date']}]")
        except Exception as e:
            print(f"ERROR: {e}")
            errors.append((f.name, str(e)))

    if not records:
        raise RuntimeError("No records extracted - check file format or xlrd installation.")

    df = pd.DataFrame(records)
    df.rename(columns=RENAME_MAP, inplace=True)

    # Default output path: inside the same folder
    if out_path is None:
        out_path = folder / f"monthly_timeseries_{folder.name}.xlsx"
    out_path = Path(out_path)

    # --- Write Excel with colour-coded headers ---
    with pd.ExcelWriter(str(out_path), engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Monthly_TimeSeries", index=False)
        ws = writer.sheets["Monthly_TimeSeries"]

        header_fills = {
            "Date": "1F3864",   # navy
            "A |":  "2E75B6",   # blue
            "F |":  "C00000",   # red
            "G |":  "375623",   # green
            "I |":  "7030A0",   # purple
        }
        thin   = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        # Format header row
        for cell in ws[1]:
            val      = str(cell.value or "")
            fill_hex = "1F3864"
            for prefix, colour in header_fills.items():
                if val.startswith(prefix):
                    fill_hex = colour
                    break
            cell.font      = Font(bold=True, color="FFFFFF", size=9)
            cell.fill      = PatternFill("solid", fgColor=fill_hex)
            cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            cell.border    = border

        # Format data rows
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border    = border
                if isinstance(cell.value, float):
                    cell.number_format = "0.00"

        # Column widths
        ws.column_dimensions["A"].width = 14
        for col_idx in range(2, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 14

        # Freeze header row
        ws.freeze_panes = "B2"

        # Error sheet if any files failed
        if errors:
            df_err = pd.DataFrame(errors, columns=["File", "Error"])
            df_err.to_excel(writer, sheet_name="Errors", index=False)

    # --- Missing dates report ---
    # Infer year and month from the dates we successfully extracted
    extracted_dates = set()
    for rec in records:
        try:
            d, m, y = rec["Date"].split("-")
            extracted_dates.add(date(int(y), int(m), int(d)))
        except Exception:
            pass

    missing_summary = ""
    if extracted_dates:
        min_date   = min(extracted_dates)
        year, month = min_date.year, min_date.month
        total_days  = calendar.monthrange(year, month)[1]

        # All calendar days for that month
        all_days = {date(year, month, d) for d in range(1, total_days + 1)}
        missing  = sorted(all_days - extracted_dates)

        missing_summary = (
            f"  Month     : {min_date.strftime('%B %Y')}  "
            f"({total_days} days in month)\n"
            f"  Extracted : {len(extracted_dates)} days\n"
        )
        if missing:
            missing_dates_str = ", ".join(d.strftime("%d-%m-%Y") for d in missing)
            missing_summary += (
                f"  Missing   : {len(missing)} day(s) -> {missing_dates_str}\n"
                f"     (files not yet published or not downloaded)\n"
            )
        else:
            missing_summary += f"  Complete  : all {total_days} days present\n"

    print(f"\n{'='*65}")
    print(missing_summary, end="")
    if errors:
        print(f"  Failed    : {len(errors)} files  (see 'Errors' sheet)")
    print(f"  Saved     : {out_path}")
    print(f"{'='*65}\n")

    return df
